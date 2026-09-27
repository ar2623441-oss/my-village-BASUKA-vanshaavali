# -*- coding: utf-8 -*-
"""
regenerate_excel.py  (v2 - self-healing + validated)
------------------------------------------------------
Rebuilds Basuka_Vanshaavali_MASTER.xlsx directly from data/basuka_vanshaavali.json.

New in this version:
  - Validates the JSON before writing anything (duplicate IDs, broken parent
    links, self-parenting, etc.) and refuses to write a broken file.
  - Auto-computes "generation" and "branch" from the parent chain if a person
    is missing them - so when you add a brand-new person you only need to set
    id / nameHi / nameEn / fatherId and the rest is derived automatically.
  - Prints a summary you can sanity-check before opening the spreadsheet.
Usage:
    pip install openpyxl --break-system-packages   (first time only)
    python regenerate_excel.py

Run it from the repo's root folder (same folder as index.html and data/).
"""
import json
import sys
import collections
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is not installed. Run:  pip install openpyxl --break-system-packages")

ROOT = Path(__file__).parent
JSON_PATH = ROOT / "data" / "basuka_vanshaavali.json"
XLSX_PATH = ROOT / "Basuka_Vanshaavali_MASTER.xlsx"

if not JSON_PATH.exists():
    sys.exit(f"Could not find {JSON_PATH}. Run this script from the repo's root folder.")

with open(JSON_PATH, encoding="utf-8") as f:
    doc = json.load(f)

people = doc["people"]

# ======================================================
# VALIDATION - catch mistakes before writing anything
# ======================================================
errors = []
warnings = []

ids_seen = collections.Counter(p["id"] for p in people)
for pid, count in ids_seen.items():
    if count > 1:
        errors.append(f"Duplicate id used {count} times: {pid}")

byid = {p["id"]: p for p in people}

for p in people:
    if p.get("fatherId"):
        if p["fatherId"] == p["id"]:
            errors.append(f"{p['id']} ({p.get('nameHi','?')}) lists itself as its own father")
        elif p["fatherId"] not in byid:
            errors.append(f"{p['id']} ({p.get('nameHi','?')}) has fatherId '{p['fatherId']}' which doesn't exist")
    if not p.get("nameHi"):
        warnings.append(f"{p['id']} has no nameHi")
    if not p.get("nameEn"):
        warnings.append(f"{p['id']} has no nameEn")

roots = [p for p in people if not p.get("fatherId")]
if len(roots) == 0:
    errors.append("No root person found (everyone has a fatherId) - the tree has no starting point")
elif len(roots) > 1:
    warnings.append(f"{len(roots)} people have no fatherId (multiple disconnected trees): "
                     + ", ".join(p["id"] for p in roots[:10]) + (" ..." if len(roots) > 10 else ""))

# cycle check (walk every node's ancestor chain, bail if it loops)
for p in people:
    seen = set()
    cur = p
    steps = 0
    while cur and cur.get("fatherId"):
        if cur["id"] in seen:
            errors.append(f"Circular ancestry detected involving {p['id']} ({p.get('nameHi','?')})")
            break
        seen.add(cur["id"])
        cur = byid.get(cur["fatherId"])
        steps += 1
        if steps > 200:
            errors.append(f"Ancestor chain for {p['id']} is suspiciously long (200+) - likely a cycle")
            break

if errors:
    print("VALIDATION FAILED - fix these in data/basuka_vanshaavali.json before re-running:\n")
    for e in errors:
        print("  -", e)
    sys.exit(1)

if warnings:
    print("Warnings (not fatal, but check these):")
    for w in warnings:
        print("  -", w)
    print()

# ======================================================
# AUTO-FILL missing generation / branch / lineEnds / page / notes
# ======================================================
BRANCH_ROOTS = {
    "P0057": "राजशाह वंश — पूरब पट्टी (Raj Shah / Poorab Patti)",
    "P0058": "भोजशाह वंश — पश्चिम पट्टी (Bhoj Shah / Paschim Patti)",
    "P0059": "विक्रमशाह वंश (Vikram Shah)",
    "P0050": "सकरवार पठान शाखा (Sakarwar Pathan)",
    "P0052": "सहजमल राय शाखा (Sahajmal Rai)",
    "P0053": "तेजमल राय शाखा (Tejmal Rai)",
    "P0054": "जटाम राय शाखा (Jatam Rai)",
    "P0055": "ठकुराई राय शाखा (Thakurai Rai)",
    "P0056": "हिन्नू राय शाखा (Hinnu Rai)",
    "P0046": "राजमल राय शाखा (Rajmal Rai)",
    "P0047": "गोसाईंमल राय शाखा (Gosaimal Rai)",
    "P0049": "संसारमल राय शाखा (Sansarmal Rai)",
    "P0031": "सैनूमल राव शाखा (Sainumal Rao)",
}

def compute_generation(p):
    if p.get("generation"):
        return p["generation"]
    depth, cur = 1, p
    while cur.get("fatherId"):
        cur = byid[cur["fatherId"]]
        depth += 1
    return depth

def compute_branch(p):
    if p.get("branch"):
        return p["branch"]
    cur = p
    while cur:
        if cur["id"] in BRANCH_ROOTS:
            return BRANCH_ROOTS[cur["id"]]
        cur = byid.get(cur["fatherId"]) if cur.get("fatherId") else None
    return "मूल वंश — ओंकार से पुरनमल राय तक (Root Lineage)"

filled_count = 0
for p in people:
    before = (p.get("generation"), p.get("branch"))
    p["generation"] = compute_generation(p)
    p["branch"] = compute_branch(p)
    p.setdefault("lineEnds", False)
    p.setdefault("page", "")
    p.setdefault("notes", "")
    p.setdefault("childrenIds", [])
    if before != (p.get("generation"), p.get("branch")):
        filled_count += 1

if filled_count:
    print(f"Auto-filled generation/branch for {filled_count} people that were missing them.\n")

# recompute childrenIds fresh from fatherId links (source of truth), ignore whatever was stored
children_names = collections.defaultdict(list)
children_ids = collections.defaultdict(list)
for p in people:
    if p.get("fatherId"):
        children_ids[p["fatherId"]].append(p["id"])
        children_names[p["fatherId"]].append(p["nameHi"])
for p in people:
    p["childrenIds"] = children_ids.get(p["id"], [])

def father_name(p, field):
    if not p.get("fatherId") or p["fatherId"] not in byid:
        return ""
    return byid[p["fatherId"]][field]

def lineage_path(p):
    chain = []
    cur = p
    while cur:
        chain.append(cur["nameHi"])
        cur = byid.get(cur["fatherId"]) if cur.get("fatherId") else None
    return " → ".join(reversed(chain))

max_generation = max(p["generation"] for p in people)

# ======================================================
# WRITE EXCEL
# ======================================================
HDR_FILL = PatternFill("solid", fgColor="7B3F00")
HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
BODY = Font(name="Arial", size=10)
THIN = Side(style="thin", color="D9CBB5")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ALT = PatternFill("solid", fgColor="FBF7F0")
ENDF = PatternFill("solid", fgColor="F6E3E3")

wb = Workbook()

# ---------- Sheet 1: Master Data ----------
ws = wb.active
ws.title = "Master Data"
HEAD = ["Person_ID","Name_Hindi","Name_English","Father_ID","Father_Name_Hindi","Father_Name_English",
        "Children_IDs","Children_Names_Hindi","Children_Count","Generation","Branch",
        "Line_Ends_Nirvansh","Is_Living_Line","PDF_Page","Lineage_Path_Hindi","Notes"]
ws.append(HEAD)
for c in ws[1]:
    c.fill = HDR_FILL; c.font = HDR_FONT
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[1].height = 34

for i, p in enumerate(people, start=2):
    kids = p["childrenIds"]
    row = [
        p["id"], p["nameHi"], p["nameEn"],
        p.get("fatherId") or "",
        father_name(p, "nameHi"),
        father_name(p, "nameEn"),
        ", ".join(kids),
        ", ".join(children_names.get(p["id"], [])),
        len(kids),
        p["generation"],
        p["branch"],
        "हाँ (×)" if p["lineEnds"] else "",
        "नहीं" if (p["lineEnds"] or not kids) else "हाँ",
        p["page"],
        lineage_path(p),
        p.get("notes", ""),
    ]
    ws.append(row)
    for c in ws[i]:
        c.font = BODY; c.border = BOX
        c.alignment = Alignment(vertical="top", wrap_text=False)
    if p["lineEnds"]:
        for c in ws[i]: c.fill = ENDF
    elif i % 2 == 0:
        for c in ws[i]: c.fill = ALT

widths = [11,26,26,11,24,24,26,42,7,11,44,14,13,10,80,52]
for idx, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(idx)].width = w
ws.freeze_panes = "C2"
ws.auto_filter.ref = "A1:P%d" % (len(people) + 1)

# ---------- Sheet 2: Branch Summary ----------
ws2 = wb.create_sheet("Branch Summary")
ws2.append(["Branch", "Persons", "Max Generation", "Lines Ended (×)"])
agg = collections.defaultdict(lambda: [0, 0, 0])
for p in people:
    a = agg[p["branch"]]
    a[0] += 1
    a[1] = max(a[1], p["generation"])
    if p["lineEnds"]:
        a[2] += 1
for branch, a in sorted(agg.items(), key=lambda kv: -kv[1][0]):
    ws2.append([branch, a[0], a[1], a[2]])
ws2.append(["TOTAL", len(people), max_generation, sum(1 for p in people if p["lineEnds"])])
for c in ws2[1]:
    c.fill = HDR_FILL; c.font = HDR_FONT
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for row in ws2.iter_rows(min_row=2):
    for c in row:
        c.font = BODY; c.border = BOX
ws2["A%d" % ws2.max_row].font = Font(name="Arial", size=10, bold=True)
for idx, w in zip(range(1, 5), [50, 12, 18, 18]):
    ws2.column_dimensions[get_column_letter(idx)].width = w

# ---------- Sheet 3: Legend & Notes ----------
ws3 = wb.create_sheet("Legend & Notes")
LEG = [
    ["बसुका ग्राम वंशावली — मास्टर डेटा / BASUKA VILLAGE VANSHAAVALI — MASTER DATA", ""],
    ["", ""],
    ["स्रोत / Source", "यह फ़ाइल data/basuka_vanshaavali.json से स्वचालित रूप से जनरेट की गई है — regenerate_excel.py चलाकर।"],
    ["कुल व्यक्ति / Total persons", len(people)],
    ["अधिकतम पीढ़ी / Max generation", max_generation],
    ["", ""],
    ["नोट / NOTE", "यह फ़ाइल केवल संदर्भ के लिए है। वेबसाइट केवल JSON फ़ाइल का उपयोग करती है — इसे संपादित करने के लिए हमेशा data/basuka_vanshaavali.json को ही सम्पादित करें, फिर इस स्क्रिप्ट को दोबारा चलाएँ।"],
    ["", "This file is for reference only. The live website reads only the JSON file - always edit data/basuka_vanshaavali.json, then re-run this script to refresh this spreadsheet."],
]
for r in LEG:
    ws3.append(r)
ws3["A1"].font = Font(name="Arial", size=14, bold=True, color="7B3F00")
for row in ws3.iter_rows(min_row=2):
    for c in row:
        if c.font.size != 14:
            c.font = Font(name="Arial", size=10)
        c.alignment = Alignment(vertical="top", wrap_text=True)
ws3.column_dimensions["A"].width = 30
ws3.column_dimensions["B"].width = 110

wb.save(XLSX_PATH)
print(f"Done - wrote {len(people)} people ({max_generation} generations) to {XLSX_PATH}")
